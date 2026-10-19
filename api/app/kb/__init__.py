"""知识库：PDF/Word/文本 → 按章、条切片（保留条款号）→ 检索（字二元组 BM25，可叠加向量）。"""
from __future__ import annotations

import math
import re
import subprocess
from collections import Counter
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import KbChunk, KbDoc

CN_NUM = "零一二三四五六七八九十百千0-9０-９"
CHAPTER_RE = re.compile(rf"^\s*(第[{CN_NUM}]+[章节部分])\s*(.*)$")
CLAUSE_RE = re.compile(rf"^\s*(第[{CN_NUM}]+条)\s*(.*)$")
NUM_RE = re.compile(r"^\s*((?:\d+\.){1,3}\d*|[一二三四五六七八九十]+、|（[一二三四五六七八九十]+）)\s*(.+)$")


def read_text(path: Path) -> str:
    suf = path.suffix.lower()
    if suf == ".pdf":
        import pymupdf
        with pymupdf.open(str(path)) as d:
            return "\n".join(p.get_text() for p in d)
    if suf == ".doc":
        out = path.with_suffix(".docx")
        subprocess.run(["soffice", "--headless", "--convert-to", "docx", "--outdir", str(path.parent), str(path)],
                       check=True, timeout=120)
        path, suf = out, ".docx"
    if suf == ".docx":
        import docx
        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    return path.read_text(encoding="utf-8", errors="ignore")


def chunk_text(text: str, max_len: int = 420) -> list[tuple[str, str]]:
    """返回 [(条款号, 文本)]。条款号形如"第二章 第十条"或"3.2"。"""
    chunks: list[tuple[str, str]] = []
    chapter = ""
    sub = ""  # 二级标题：（一）身心状况 / 关键指标 B5.生活照料
    cur_no, cur = "", []

    def flush():
        body = "\n".join(cur).strip()
        if body:
            label = " ".join(x for x in (chapter, sub, cur_no) if x).strip() or "正文"
            while len(body) > max_len:
                cut = body.rfind("。", 0, max_len)
                cut = cut + 1 if cut > max_len // 3 else max_len
                chunks.append((label, body[:cut]))
                body = body[cut:].strip()
            if body:
                chunks.append((label, body))

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        m = CHAPTER_RE.match(line)
        if m and len(line) < 40:
            flush(); cur, cur_no = [], ""
            chapter = (m.group(1) + " " + m.group(2)).strip()
            continue
        if re.match(r"^[一二三四五六七八九十]+、", line) and len(line) <= 16:  # "一、入园与离园"这类小标题视为章
            flush(); cur, cur_no, sub = [], "", ""
            chapter = line
            continue
        if line.startswith("重点内容 ") and len(line) <= 24:  # 评估指南：重点内容 A2.保育与安全
            flush(); cur, cur_no, sub = [], "", ""
            chapter = line.replace("重点内容 ", "")
            continue
        if line.startswith("维度：") and len(line) <= 20:  # 专业标准：维度：专业理念与师德
            flush(); cur, cur_no, sub = [], "", ""
            chapter = line.replace("维度：", "")
            continue
        if (line.startswith("关键指标 ") and len(line) <= 24) or (re.match(r"^[（(][一二三四五六七八九十]+[）)]", line) and len(line) <= 16):
            flush(); cur, cur_no = [], ""
            sub = re.sub(r"^\(([一二三四五六七八九十]+)\)", lambda m: "（" + m.group(1) + "）", line.replace("关键指标 ", ""))
            continue
        if line.startswith("#"):
            continue
        m = re.match(r"^(目标\d+)\s*(.*)$", line) or re.match(r"^(考查要点\d+)[.．]\s*(.*)$", line)
        if m:  # 指南"目标N"、评估指南"考查要点N"作为最小引用单位
            flush()
            cur_no = m.group(1)
            cur = [line]
            continue
        if line.startswith("教育建议") and cur_no.startswith("目标"):
            cur.append(line)
            continue
        m = CLAUSE_RE.match(line) or NUM_RE.match(line)
        if m:
            flush()
            cur_no = m.group(1).rstrip("、.")
            if re.fullmatch(r"\d+", cur_no):
                cur_no = f"第{cur_no}条"
            elif re.fullmatch(r"[一二三四五六七八九十]+", cur_no):
                cur_no = f"{cur_no}、"
            cur = [line]
            continue
        cur.append(line)
    flush()
    return chunks


def _grams(s: str) -> list[str]:
    s = re.sub(r"[\s，。、；：！？“”‘’（）()《》\[\]【】,.;:!?\-—/|]+", "", s)
    return [s[i:i + 2] for i in range(len(s) - 1)] + list(s)


class _Index:
    def __init__(self):
        self.version = -1
        self.docs: list[tuple[int, Counter, int]] = []
        self.df: Counter = Counter()
        self.avg = 1.0


_idx: dict[str, _Index] = {}


def _index(db: Session, visibility: str | None) -> _Index:
    key = visibility or "all"
    total = db.scalar(select(KbChunk.id).order_by(KbChunk.id.desc()).limit(1)) or 0
    cnt = len(db.scalars(select(KbChunk.id)).all())
    ver = total * 100000 + cnt
    ix = _idx.get(key)
    if ix and ix.version == ver:
        return ix
    ix = _Index()
    ix.version = ver
    q = select(KbChunk, KbDoc).join(KbDoc, KbDoc.id == KbChunk.doc_id).where(KbDoc.status == "ready")
    if visibility == "public":
        q = q.where(KbDoc.visibility == "public")
    for ch, doc in db.execute(q):
        tf = Counter(_grams(doc.title + " " + ch.clause_no + " " + ch.text))
        ix.docs.append((ch.id, tf, sum(tf.values())))
        ix.df.update(tf.keys())
    ix.avg = (sum(d[2] for d in ix.docs) / len(ix.docs)) if ix.docs else 1.0
    _idx[key] = ix
    return ix


def search(db: Session, query: str, k: int = 8, visibility: str | None = None) -> list[dict]:
    ix = _index(db, visibility)
    if not ix.docs:
        return []
    qg = Counter(_grams(query))
    n = len(ix.docs)
    scored = []
    for cid, tf, ln in ix.docs:
        s = 0.0
        for g, qc in qg.items():
            f = tf.get(g)
            if not f:
                continue
            idf = math.log(1 + (n - ix.df[g] + 0.5) / (ix.df[g] + 0.5))
            w = 1.6 if len(g) == 2 else 0.25
            s += w * idf * (f * 2.2) / (f + 1.2 * (0.25 + 0.75 * ln / ix.avg))
        if s > 0:
            scored.append((s, cid))
    scored.sort(reverse=True)
    out = []
    for s, cid in scored[:k]:
        ch = db.get(KbChunk, cid)
        out.append({"id": ch.id, "clause_no": ch.clause_no, "text": ch.text, "doc_title": ch.doc.title,
                    "doc_kind": ch.doc.kind, "score": round(s, 3)})
    return out


def ingest(db: Session, doc: KbDoc, text: str | None = None) -> int:
    if text is None:
        text = read_text(Path(doc.file_path))
    db.query(KbChunk).filter(KbChunk.doc_id == doc.id).delete()
    parts = chunk_text(text)
    for no, body in parts:
        db.add(KbChunk(doc_id=doc.id, clause_no=no, text=body))
    doc.chunk_count = len(parts)
    doc.status = "ready"
    db.flush()
    _idx.clear()
    return len(parts)
